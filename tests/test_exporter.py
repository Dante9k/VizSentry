from pathlib import Path
import hashlib

from openpyxl import load_workbook
from openpyxl.utils.escape import unescape
from PIL import Image

from vizsentry.database import Database
from vizsentry.exporter import export_workbook


def test_export_keeps_every_row_and_full_unicode_logs(tmp_path):
    db = Database(tmp_path / 'db.sqlite3')
    job = db.create_job({'url': 'https://tenant.feishu.cn/base/test'})['id']
    db.add_records(job, [{'record_id': '=HYPERLINK("bad")', 'images': [], 'log_urls': ['https://example.test/log']}, {'record_id': 'empty'}], 0)
    db.update_job(job, total=2, import_complete=1)
    text = ('日志=不执行公式\nERROR detail 🚀\r\n' * 3000) + 'END'
    log = tmp_path / 'full.txt'
    with log.open('w', encoding='utf-8', newline='') as file:
        file.write(text)
    picture = tmp_path / 'thumb.png'
    Image.new('RGB', (160, 100), 'teal').save(picture)
    db.add_artifact(job, 'pic-thumb', picture, 'image', '图.png')
    db.add_artifact(job, 'log-text', log, 'text', '完整日志')
    db.finish_row(job, '=HYPERLINK("bad")', {'images': [{'id': 'pic', 'name': '图.png', 'verdict': '正常', 'reason': '文字完整'}, {'id': 'pic2', 'name': '另一图.png', 'verdict': '异常', 'reason': '画面缺失'}], 'logs': [{'id': 'log', 'name': '日志', 'complete': True}], 'image_verdict': '待确认', 'image_reason': '=1+1', 'log_summary': '完整日志测试'})
    path = export_workbook(db, job, tmp_path / 'report.xlsx')
    book = load_workbook(path)
    sheet = book['结果']
    assert sheet.max_row == 3
    assert sheet['B2'].data_type == 's' and sheet['F2'].value == '=1+1'
    assert sheet['B3'].value == 'empty'
    assert len(sheet._images) == 1
    assert sheet['K2'].hyperlink.target == 'https://tenant.feishu.cn/base/test'
    assert book['图片详情'].max_row == 3
    assert book['图片详情']['E2'].value == '正常'
    assert book['图片详情']['F3'].value == '画面缺失'
    assert book['日志全文'].sheet_state == 'hidden'
    # Excel preserves CR using OOXML _x000D_; openpyxl leaves it escaped in inline strings.
    restored = ''.join(unescape(row[3].value) for row in list(book['日志全文'].rows)[1:])
    assert hashlib.sha256(restored.encode()).digest() == hashlib.sha256(text.encode()).digest(), ascii(restored[:100])


def test_empty_table_exports_headers_only(tmp_path):
    db = Database(tmp_path / 'db.sqlite3')
    job = db.create_job({})['id']
    db.update_job(job, import_complete=1)
    book = load_workbook(export_workbook(db, job, tmp_path / 'empty.xlsx'))
    assert book['结果'].max_row == 1
